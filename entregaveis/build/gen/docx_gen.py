# -*- coding: utf-8 -*-
from base import *
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
OUT=ROOT+"/build/out/"
def shade(p,hexc):
    pPr=p._p.get_or_add_pPr();s=OxmlElement('w:shd');s.set(qn('w:val'),'clear');s.set(qn('w:color'),'auto');s.set(qn('w:fill'),hexc);pPr.append(s)
def new(titulo,sub):
    d=Document()
    for s in d.sections: s.left_margin=s.right_margin=Cm(2.2);s.top_margin=s.bottom_margin=Cm(2)
    st=d.styles["Normal"];st.font.name="Calibri";st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    for h,sz in (("Heading 1",17),("Heading 2",13.5),("Heading 3",12)):
        d.styles[h].font.name="Calibri";d.styles[h].font.size=Pt(sz);d.styles[h].font.color.rgb=RGBColor(0x0B,0x12,0x20);d.styles[h].font.bold=True
    p=d.add_paragraph();r=p.add_run("KIT NÃO VIRE ESTATÍSTICA · MODELO EDITÁVEL");r.font.size=Pt(8.5);r.bold=True;r.font.color.rgb=RGBColor(0xB4,0x83,0x00)
    d.add_heading(titulo,1)
    p=d.add_paragraph(sub);p.runs[0].font.color.rgb=RGBColor(0x5B,0x65,0x76)
    note(d,"Como usar: troque tudo o que está entre [colchetes] pelos seus dados e apague esta caixa. Modelo de referência: para valores altos ou contratos complexos, peça revisão a um advogado ou ao seu contador.")
    ft=d.sections[0].footer.paragraphs[0];ft.text="Modelo de referência · Kit Não Vire Estatística · não substitui orientação jurídica";ft.runs[0].font.size=Pt(8);ft.runs[0].font.color.rgb=RGBColor(0x8B,0x98,0xAD)
    return d
def note(d,t):
    p=d.add_paragraph(t);shade(p,"FFF4CC");p.runs[0].font.size=Pt(9.5);return p
def para(d,t,bold=False,just=True):
    p=d.add_paragraph();r=p.add_run(t);r.bold=bold
    if just:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p
def clause(d,n,t,body):
    p=d.add_paragraph();r=p.add_run(f"CLÁUSULA {n} – {t}");r.bold=True
    for b in body: para(d,b)
def sign(d,local=True,partes=("[CONTRATANTE]","[CONTRATADA]"),testemunhas=True):
    d.add_paragraph()
    if local: para(d,"[Cidade], [dia] de [mês] de [ano].",just=False)
    for pt in partes:
        d.add_paragraph();d.add_paragraph("_____________________________________________");d.add_paragraph(pt+"\nNome / CPF ou CNPJ")
    if testemunhas:
        d.add_paragraph();para(d,"Testemunhas:",True)
        d.add_paragraph("1. ____________________________  CPF: ______________\n2. ____________________________  CPF: ______________")
def save(d,n):
    p=OUT+n;d.save(p);print(p)

CLAUSULA_TRIB=["Os preços deste contrato foram fixados considerando a legislação tributária vigente na data da sua assinatura, incluindo os tributos incidentes sobre o fornecimento (PIS/Cofins, ICMS e/ou ISS, conforme o caso).",
 "Parágrafo primeiro. A partir da entrada em vigor da Contribuição sobre Bens e Serviços (CBS) e do Imposto sobre Bens e Serviços (IBS), instituídos pela Emenda Constitucional nº 132/2023 e regulamentados pela Lei Complementar nº 214/2025, a CBS e o IBS serão destacados no documento fiscal e acrescidos ao preço líquido aqui ajustado (\"cobrança por fora\"), substituindo os tributos que deixarem de existir, nos percentuais e no cronograma de transição definidos em lei.",
 "Parágrafo segundo. Para fins deste contrato, o \"preço líquido\" de cada item corresponde ao preço vigente na data de assinatura, deduzidos os tributos sobre o consumo nele embutidos, a saber: [__]% de PIS/Cofins e [__]% de ICMS/ISS, totalizando [__]%.",
 "Parágrafo terceiro. Havendo criação, extinção ou alteração de tributos, alíquotas, bases de cálculo, regimes de crédito ou forma de recolhimento (inclusive o recolhimento na liquidação financeira, \"split payment\") que comprovadamente alterem o custo do fornecimento, as partes revisarão o preço para restabelecer o equilíbrio econômico-financeiro original, para mais ou para menos, mediante comunicação escrita com memória de cálculo e antecedência mínima de [30] dias.",
 "Parágrafo quarto. A revisão de que trata esta cláusula não se confunde com o reajuste periódico por índice inflacionário previsto na cláusula [__], e com ele não se compensa.",
 "Parágrafo quinto. Não havendo acordo sobre o novo preço em até [30] dias da comunicação, qualquer das partes poderá rescindir este contrato sem multa, mediante aviso prévio de [30] dias, sendo devidos os valores dos fornecimentos já realizados."]

def politica():
    d=new("Política de Venda a Prazo e Cobrança","Para comércio, alimentação e varejo: regras claras de quem pode comprar fiado/a prazo, quanto e como cobrar.")
    d.add_heading("1. Objetivo",2);para(d,"Esta política define as regras de venda a prazo da [NOME DA EMPRESA] para proteger o caixa, reduzir a inadimplência e tratar todos os clientes de forma igual e respeitosa. Vale a partir de [data] e é revista a cada [6] meses.")
    d.add_heading("2. Formas de pagamento aceitas",2)
    for t in ["À vista: dinheiro, Pix, cartão de débito.","Cartão de crédito: em até [3]x sem juros acima de R$ [__]; acima de [3]x, com acréscimo informado antes da compra.","A prazo (crediário próprio, boleto ou caderneta/fiado): somente para clientes cadastrados e aprovados conforme o item 3."]: d.add_paragraph(t,style="List Bullet")
    d.add_heading("3. Quem pode comprar a prazo",2)
    for t in ["Cadastro completo: nome, CPF/CNPJ, endereço, telefone/WhatsApp e e-mail, com documento conferido.","Consulta a birô de crédito (ex.: Serasa, SPC, Boa Vista) [se a empresa tiver acesso].","Pelo menos [3] compras à vista nos últimos [90] dias (cliente novo começa no limite inicial).","Sem parcela em atraso há mais de [5] dias."]: d.add_paragraph(t,style="List Bullet")
    d.add_heading("4. Limites",2)
    t=d.add_table(rows=1,cols=3);t.style="Light Grid Accent 1"
    for i,h in enumerate(["Perfil","Limite de crédito","Prazo máximo"]): t.rows[0].cells[i].text=h
    for row in [("Cliente novo aprovado","R$ [200]","[15] dias"),("Cliente com 6 meses sem atraso","R$ [500]","[30] dias"),("Cliente com 12 meses sem atraso","R$ [1.000]","[30] dias, até [2] parcelas"),("Cliente empresa (PJ)","Análise caso a caso","[28] dias, boleto")]:
        c=t.add_row().cells
        for i,v in enumerate(row): c[i].text=v
    para(d,"Teto da carteira: o total vendido a prazo e ainda não recebido não pode passar de [15]% do faturamento médio mensal. Atingido o teto, novas vendas a prazo só com autorização de [responsável].")
    d.add_heading("5. Régua de cobrança",2)
    for m in NVE["regua"]:
        p=d.add_paragraph(style="List Bullet");r=p.add_run(f'{m["quando"]} – {m["titulo"]}: ');r.bold=True;p.add_run(m["msg"].replace("{","[").replace("}","]"))
    d.add_heading("6. Atraso",2)
    for t in ["Multa de [2]% sobre o valor em atraso (limite do Código de Defesa do Consumidor, art. 52, §1º, para vendas a prazo ao consumidor) e juros de mora de [1]% ao mês, informados no momento da compra.","Atraso acima de [5] dias: bloqueio de novas compras a prazo até a regularização.","Atraso acima de [30] dias: comunicação escrita prévia e, se não houver acordo, inclusão em cadastro de inadimplentes (CDC, art. 43, §2º) e/ou protesto.","Acordos sempre por escrito (modelo de Termo de Acordo no Pack ou no app)."]: d.add_paragraph(t,style="List Bullet")
    d.add_heading("7. Regras de respeito na cobrança",2)
    para(d,"É proibido expor o cliente ao ridículo, constrangê-lo ou ameaçá-lo (CDC, art. 42); cobrar em público, com colegas de trabalho, vizinhos ou parentes; ou divulgar listas de devedores. Contatos só entre 8h e 20h em dias úteis e [9h e 14h] aos sábados.")
    d.add_heading("8. Indicadores acompanhados toda semana",2)
    for t in ["Total a receber e quanto está vencido (por faixa: 1–15, 16–30, 31–60, +60 dias).","% de inadimplência = vencido há mais de 30 dias ÷ total vendido a prazo nos últimos 90 dias. Meta: abaixo de [3]%.","Prazo médio de recebimento (dias)."]: d.add_paragraph(t,style="List Bullet")
    d.add_heading("Termo de ciência do cliente (imprimir no cadastro)",2)
    para(d,"Eu, [nome], CPF [__], declaro que conheço as condições de venda a prazo da [EMPRESA]: limite de R$ [__], vencimento em até [__] dias, multa de [2]% e juros de [1]% ao mês em caso de atraso, e autorizo o contato pelo WhatsApp/e-mail informados para lembretes de pagamento.")
    sign(d,partes=("[CLIENTE]","[EMPRESA]"),testemunhas=False)
    save(d,"Modelo-Politica-Venda-a-Prazo.docx")

def clausula_servicos():
    d=new("Cláusula de Reajuste Tributário para Contratos de Serviços","Para incluir em contratos e propostas novos e em renovações (consultorias, agências, clínicas, escolas, TI, manutenção e serviços recorrentes).")
    d.add_heading("Por que usar",2)
    para(d,"Em 2027 a CBS substitui o PIS/Cofins e, de 2029 a 2033, o IBS substitui gradualmente o ISS. Contratos longos assinados agora vão atravessar essa mudança. Sem uma cláusula, mudar o preço depende só da boa vontade do cliente. Com ela, a revisão é um direito das duas partes, para mais ou para menos, com regra clara.")
    d.add_heading("Versão A · Cláusula completa (contratos acima de 6 meses)",2)
    clause(d,"[__]","DOS TRIBUTOS E DA REVISÃO POR ALTERAÇÃO TRIBUTÁRIA",CLAUSULA_TRIB)
    d.add_heading("Versão B · Cláusula curta (propostas e contratos simples)",2)
    para(d,"Os valores desta proposta consideram os tributos vigentes em [data]. A partir de 2027, a CBS e o IBS (EC 132/2023 e LC 214/2025) serão destacados na nota fiscal e somados ao valor líquido do serviço, em substituição ao PIS/Cofins e ao ISS, conforme o cronograma legal. Qualquer outra alteração tributária que mude o custo do serviço permitirá a revisão do preço, para mais ou para menos, mediante aviso com [30] dias de antecedência.")
    d.add_heading("Versão C · Texto para rodapé de orçamento",2)
    para(d,"Valores com tributos vigentes em [data]. Sujeitos a revisão conforme a Reforma Tributária (CBS/IBS a partir de 2027).")
    d.add_heading("Carta de comunicação de revisão (quando aplicar a cláusula)",2)
    para(d,"[Cidade], [data].\n\nÀ [CLIENTE]\nRef.: Contrato nº [__] – revisão por alteração tributária\n\nPrezados,\n\nConforme a cláusula [__] do contrato, comunicamos que, a partir de [data], a [alteração: ex. substituição do PIS/Cofins pela CBS] altera o custo tributário do serviço. Segue a memória de cálculo:\n\n• Preço atual: R$ [__]\n• Tributos embutidos hoje: [__]% (R$ [__])\n• Preço líquido: R$ [__]\n• Tributos a partir de [data]: [__]% destacados na nota (R$ [__])\n• Novo valor total da nota: R$ [__]\n\nLembramos que, sendo a sua empresa do regime regular, a CBS/IBS destacada na nota poderá gerar crédito para vocês, conforme a legislação. Ficamos à disposição para conversar até [data].\n\nAtenciosamente,\n[nome] – [empresa] – [telefone]",just=False)
    note(d,"Dica: rode o simulador Preço 2027 do app (ou a aba Preço 2027 do Painel) para preencher a memória de cálculo. Peça ao contador para conferir os percentuais.")
    save(d,"Modelo-Clausula-Reajuste-Tributario-Contrato.docx")

def aditivo():
    d=new("Termo Aditivo Contratual – Cláusula Tributária","Para contratos de fornecimento já assinados (indústria, agro, distribuição e B2B) que vão atravessar a transição 2027–2033.")
    para(d,"TERMO ADITIVO Nº [__] AO CONTRATO DE [FORNECIMENTO / PRESTAÇÃO DE SERVIÇOS] Nº [__], celebrado em [data]",True)
    para(d,"Pelo presente instrumento, de um lado [RAZÃO SOCIAL], CNPJ [__], com sede em [endereço], neste ato representada por [nome, cargo], doravante CONTRATADA/FORNECEDORA; e de outro [RAZÃO SOCIAL], CNPJ [__], com sede em [endereço], neste ato representada por [nome, cargo], doravante CONTRATANTE/COMPRADORA; resolvem celebrar o presente Termo Aditivo, mediante as cláusulas seguintes:")
    clause(d,"PRIMEIRA","DO OBJETO",["O presente aditivo tem por objeto adequar o Contrato à Reforma Tributária do consumo (Emenda Constitucional nº 132/2023 e Lei Complementar nº 214/2025), estabelecendo regras de formação de preço, destaque de tributos, créditos e revisão durante o período de transição e após a sua conclusão."])
    clause(d,"SEGUNDA","DO PREÇO LÍQUIDO E DOS TRIBUTOS",CLAUSULA_TRIB)
    clause(d,"TERCEIRA","DOS CRÉDITOS E DOCUMENTOS FISCAIS",["A FORNECEDORA compromete-se a emitir documento fiscal idôneo, com destaque correto da CBS e do IBS, e a manter-se regular perante o fisco, de modo a viabilizar a apropriação de créditos pela COMPRADORA nos termos da LC 214/2025.",
      "Parágrafo único. Caso a FORNECEDORA altere o seu regime tributário (por exemplo, opção ou desistência do recolhimento de IBS/CBS pelo regime regular, sendo optante do Simples Nacional) de forma que modifique o crédito apropriável pela COMPRADORA, deverá comunicá-lo por escrito com antecedência mínima de [60] dias, facultada às partes a revisão do preço nos termos da Cláusula Segunda."])
    clause(d,"QUARTA","DO RECOLHIMENTO NA LIQUIDAÇÃO FINANCEIRA (SPLIT PAYMENT)",["As partes reconhecem que, quando implantado o recolhimento da CBS e do IBS na liquidação financeira, parte do valor pago poderá ser segregada e recolhida diretamente ao fisco. Os prazos de pagamento e as formas de liquidação deste Contrato serão mantidos, e eventuais ajustes operacionais (meios de pagamento, boletos, conciliação) serão acordados por escrito sem alteração do preço líquido."])
    clause(d,"QUINTA","DO PRAZO DE TRANSIÇÃO",["As regras deste aditivo aplicam-se a todos os fornecimentos cujo fato gerador ocorra a partir de [01/01/2027], inclusive nos anos de transição (2027 a 2032) e após a extinção do ICMS e do ISS (2033)."])
    clause(d,"SEXTA","DA RATIFICAÇÃO",["Permanecem inalteradas e em pleno vigor as demais cláusulas do Contrato que não conflitem com este aditivo.","E, por estarem justas e acordadas, as partes assinam o presente em [2] vias de igual teor, ou eletronicamente."])
    sign(d,partes=("[CONTRATANTE / COMPRADORA]","[CONTRATADA / FORNECEDORA]"))
    d.add_page_break();d.add_heading("Anexo I – Memória de preço líquido",2)
    t=d.add_table(rows=1,cols=6);t.style="Light Grid Accent 1"
    for i,h in enumerate(["Item / código (NCM)","Preço atual","Tributos embutidos %","Preço líquido","CBS/IBS a partir de 2027 %","Preço com CBS/IBS"]): t.rows[0].cells[i].text=h
    for _ in range(8): t.add_row()
    note(d,"Dica: use o simulador Preço 2027 do app ou a aba Preço 2027 do Painel para preencher este anexo.")
    save(d,"Modelo-Aditivo-Contratual-Clausula-Tributaria.docx")

def pack():
    d=new("Pack de Modelos Editáveis – Renegociação e Cobrança","Termo de acordo e confissão de dívida, proposta ao fornecedor, carta ao banco e todas as mensagens do Pack.")
    d.add_heading("1. Termo de Acordo e Confissão de Dívida com Parcelamento",2)
    para(d,"CREDOR: [RAZÃO SOCIAL/NOME], CNPJ/CPF [__], endereço [__].\nDEVEDOR: [NOME], CPF/CNPJ [__], endereço [__], telefone [__], e-mail [__].",just=False)
    clause(d,"1ª","DA DÍVIDA",["O DEVEDOR reconhece e confessa dever ao CREDOR a quantia líquida, certa e exigível de R$ [valor] ([por extenso]), referente a [descrição: notas fiscais nº __, compras de __ a __, serviços de __], vencida(s) em [data(s)]."])
    clause(d,"2ª","DO PARCELAMENTO",["O valor será pago da seguinte forma: entrada de R$ [__] em [data] e [n] parcelas mensais e sucessivas de R$ [__], vencendo a primeira em [data] e as demais no mesmo dia dos meses seguintes, por Pix (chave [__]) / boleto / depósito na conta [__].","Parágrafo único. [Opcional: Sobre o saldo incidem juros de [__]% ao mês, já incluídos nas parcelas acima.]"])
    clause(d,"3ª","DO ATRASO E DO VENCIMENTO ANTECIPADO",["O atraso de qualquer parcela por mais de [10] dias implica o vencimento antecipado das demais, acrescidas de multa de [2]% e juros de mora de [1]% ao mês, podendo o CREDOR adotar as medidas de cobrança cabíveis, inclusive protesto e inclusão em cadastro de inadimplentes, mediante comunicação prévia."])
    clause(d,"4ª","DA QUITAÇÃO",["Pago integralmente o acordo, o CREDOR dará plena quitação da dívida descrita na Cláusula 1ª e, se for o caso, providenciará a baixa de protestos ou registros em até [5] dias úteis."])
    clause(d,"5ª","DISPOSIÇÕES GERAIS",["Este termo, assinado pelo devedor e por duas testemunhas, constitui título executivo extrajudicial (Código de Processo Civil, art. 784, III). Fica eleito o foro da comarca de [__]."])
    sign(d,partes=("[CREDOR]","[DEVEDOR]"))
    d.add_page_break();d.add_heading("2. Proposta de Renegociação para Fornecedor",2)
    para(d,"[Cidade], [data].\n\nÀ [FORNECEDOR]\nA/C [nome do contato / financeiro]\nRef.: Proposta de regularização e condições comerciais\n\nPrezados,\n\nSomos clientes da [FORNECEDOR] desde [ano], com compras médias de R$ [__] por mês. Estamos reorganizando o nosso caixa para 2027 e queremos manter a parceria em dia e transparente.\n\nHoje temos em aberto R$ [__] (notas [__]). Propomos:\n\n1. Entrada de R$ [__] em [data];\n2. Saldo de R$ [__] em [n] parcelas de R$ [__], a partir de [data];\n3. Novas compras durante o acordo: [à vista / prazo de __ dias];\n4. [Opcional] Prazo das compras a partir de [mês] passa de [28] para [42] dias, com pedido programado mensal de R$ [__].\n\nPedimos retorno até [data]. Estamos à disposição para ajustar a proposta em uma conversa.\n\nAtenciosamente,\n[nome] – [cargo] – [empresa] – [telefone]",just=False)
    d.add_page_break();d.add_heading("3. Carta de Proposta de Renegociação ao Banco",2)
    para(d,"[Cidade], [data].\n\nAo [BANCO] – Agência [__]\nA/C Gerente [nome]\nRef.: Pedido de renegociação e troca de dívidas – [RAZÃO SOCIAL], CNPJ [__], conta [__]\n\nPrezado(a) [nome],\n\nSomos clientes desde [ano] e queremos reorganizar o nosso endividamento antes de qualquer atraso. A posição atual é:\n\n• [Cheque especial / conta garantida]: saldo R$ [__], taxa [__]% a.m.\n• [Capital de giro]: saldo R$ [__], taxa [__]% a.m., parcela R$ [__], [n] parcelas restantes.\n• [Cartão / antecipação]: saldo R$ [__], taxa [__]% a.m.\n\nTotal: R$ [__]. Faturamento médio dos últimos 12 meses: R$ [__] por mês (fluxo de caixa de 13 semanas em anexo).\n\nSolicitamos:\n1. Consolidação dos saldos em uma única operação com taxa menor e prazo de [__] meses;\n2. Carência de [__] meses para a primeira parcela;\n3. Avaliação de linhas com garantia do FGO/Pronampe ou de fundos garantidores, se elegíveis;\n4. Parcela máxima compatível com o nosso caixa: R$ [__].\n\nPedimos a apresentação das propostas com o Custo Efetivo Total (CET) mensal e anual, por escrito, até [data].\n\nAtenciosamente,\n[nome] – [cargo] – [telefone]\n\nAnexos: lista de dívidas; fluxo de caixa 13 semanas; faturamento 12 meses.",just=False)
    d.add_page_break();d.add_heading("4. Régua de cobrança (WhatsApp/e-mail)",2)
    for m in NVE["regua"]:
        d.add_heading(f'{m["quando"]} – {m["titulo"]}',3);p=para(d,m["msg"].replace("{","[").replace("}","]"),just=False);shade(p,"F1F4F8")
    d.add_heading("5. Mensagem de acordo fechado",3)
    p=para(d,'Olá, [nome]! Conforme combinamos hoje, o débito de [valor total] será pago em [n] parcelas de [valor], com vencimentos em [datas], pelo Pix [chave]. Se alguma parcela atrasar mais de [5] dias, o acordo é cancelado e o saldo volta a valer com multa e juros previstos. Pode me confirmar com um "de acordo"? Obrigado! [empresa]',just=False);shade(p,"F1F4F8")
    d.add_heading("6. E-mail formal para cliente empresa (B2B)",3)
    p=para(d,"Assunto: Pendência financeira – NF [número] – [sua empresa]\n\nPrezados, identificamos em aberto a nota fiscal [número], no valor de [valor], vencida em [data]. Solicitamos a regularização até [data] ou o contato para alinharmos uma proposta. Seguem os dados para pagamento: [dados]. Caso o pagamento já tenha sido feito, pedimos a gentileza de enviar o comprovante.\n\nAtenciosamente, [nome] – [cargo] – [telefone]",just=False);shade(p,"F1F4F8")
    d.add_heading("7. Falas para fornecedores e banco",3)
    for t in ['Mais prazo: "[Nome], somos clientes há [tempo] e compramos em média [valor] por mês com vocês. Estamos reorganizando o caixa para 2027 e queria propor passar o prazo de [28] para [42] dias nas próximas compras. Em troca, consigo concentrar mais pedidos com vocês. Isso é possível?"',
              'Desconto à vista: "Se eu antecipar o pagamento desta compra para o Pix hoje, qual desconto vocês conseguem? Pela minha conta, [3%] faz sentido para os dois lados."',
              'Banco: "Quero trocar esse saldo por uma linha mais barata e mais longa. Com uma parcela de até [valor], eu pago em dia. Quais opções você consegue me apresentar até [data]? Preciso do CET por escrito para comparar."',
              'Se o banco disser não: "Entendo. Pode registrar o meu pedido e me dizer o que precisaria mudar para aprovar? Vou cotar em outras instituições e volto a falar com você."']:
        p=para(d,t,just=False);shade(p,"F1F4F8")
    save(d,"Pack-Modelos-Editaveis.docx")
if __name__=="__main__":
    politica();clausula_servicos();aditivo();pack()
