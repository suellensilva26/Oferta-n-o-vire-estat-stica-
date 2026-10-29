# -*- coding: utf-8 -*-
from base import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference
from datetime import date
OUT=ROOT+"/build/out/"
DARK=PatternFill("solid",fgColor="0B1220");GOLD=PatternFill("solid",fgColor="FCC419");INP=PatternFill("solid",fgColor="FFF4CC");RES=PatternFill("solid",fgColor="E8F7EE");GREY=PatternFill("solid",fgColor="F1F4F8")
RED=PatternFill("solid",fgColor="FDE2E2");GRN=PatternFill("solid",fgColor="DCFCE7")
thin=Side(style="thin",color="D5DBE5");BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
BRL='"R$" #,##0.00;[Red]-"R$" #,##0.00';BRL0='"R$" #,##0;[Red]-"R$" #,##0';PCT='0.00%';PCT1='0.0%'
def title(ws,t,sub=None,width=6):
    ws.sheet_view.showGridLines=False
    ws["A1"]=t;ws["A1"].font=Font(name="Montserrat",bold=True,size=16,color="FFFFFF")
    for c in range(1,width+1): ws.cell(1,c).fill=DARK; ws.cell(2,c).fill=DARK
    ws.row_dimensions[1].height=30
    if sub: ws["A2"]=sub;ws["A2"].font=Font(name="Inter",size=10,color="FCC419")
def hdr(ws,row,vals,col=1):
    for i,v in enumerate(vals):
        c=ws.cell(row,col+i,v);c.fill=DARK;c.font=Font(bold=True,color="FFFFFF",name="Inter");c.alignment=Alignment(wrap_text=True,vertical="center");c.border=BOX
def inp(c,fmt=None):
    c.fill=INP;c.border=BOX;c.font=Font(name="Inter",bold=True,color="0B1220")
    if fmt:c.number_format=fmt
def res(c,fmt=None,bold=True):
    c.fill=RES;c.border=BOX;c.font=Font(name="Inter",bold=bold)
    if fmt:c.number_format=fmt
def lab(c,bold=False):
    c.font=Font(name="Inter",bold=bold);c.alignment=Alignment(wrap_text=True,vertical="center");c.border=BOX
def widths(ws,d):
    for k,v in d.items(): ws.column_dimensions[k].width=v
def leia(wb,titulo,linhas):
    ws=wb.create_sheet("Leia-me",0);title(ws,titulo,"Kit Não Vire Estatística · Plano de Blindagem 2027–2030",3)
    r=4
    for l in linhas:
        ws.cell(r,1,l).alignment=Alignment(wrap_text=True,vertical="top");ws.cell(r,1).font=Font(name="Inter",size=11,bold=l.startswith("•")==False and l.endswith(":"))
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3);ws.row_dimensions[r].height=max(18,15*(len(l)//95+1));r+=1
    widths(ws,{"A":45,"B":45,"C":45})
    c=ws.cell(r+1,1,"Legenda:");c.font=Font(bold=True)
    inp(ws.cell(r+2,1,"Célula amarela = você preenche"));res(ws.cell(r+3,1,"Célula verde = calculada (não apague)"),bold=False)
    ws.cell(r+5,1,"Conteúdo educativo; não substitui contador. Alíquotas de 2027+ são estimativas não oficiais (dados de 28/10/2026).").font=Font(italic=True,size=9,color="5B6576")

# ---------------- CALCULADORA ----------------
def calculadora():
    wb=Workbook();wb.remove(wb.active)
    leia(wb,"Calculadora de Impacto da Reforma",["Para que serve:","• Estimar quanto a troca de PIS/Cofins/ICMS/ISS por CBS/IBS muda o imposto, o lucro e o preço do seu negócio, ano a ano, de 2026 a 2033.",
     "Como usar:","• Aba 1: preencha as células amarelas com números de um mês típico.","• Aba 2: as alíquotas são EDITÁVEIS. Quando o Senado fixar a alíquota oficial da CBS (até 15/12/2026), troque a célula B4. Tudo recalcula.",
     "• Aba 3: veja a carga de impostos ano a ano e o reajuste de preço necessário para manter o lucro de hoje.","• Aba 4 (só Simples): compare ficar 'por dentro' x ir para o regime regular de IBS/CBS ('híbrido').",
     "Funciona no Excel e no Google Planilhas (drive.google.com → Novo → Upload → Abrir com Planilhas Google).",
     "Limites: não simula regimes específicos (combustíveis, financeiro, imóveis, cooperativas, ZFM), substituição tributária nem benefícios de ICMS. Considera que os fornecedores do regime regular repassam IBS/CBS no preço."])
    ws=wb.create_sheet("1. Seus dados");title(ws,"1. Seus dados","Preencha as células amarelas com valores de um MÊS típico",3)
    rows=[("Faturamento bruto do mês",40000,BRL,"Tudo o que vendeu no mês, com impostos."),
     ("Compras com nota fiscal no mês (mercadoria, insumos, serviços, energia, frete)",18000,BRL,"Só compras com nota. Salário não entra."),
     ("% dessas compras de fornecedores do regime regular (Presumido/Real ou Simples híbrido)",0.6,PCT1,"Compra de MEI, Simples 'por dentro' ou pessoa física gera pouco ou nenhum crédito."),
     ("Folha + encargos + pró-labore do mês",9000,BRL,"Não gera crédito de IBS/CBS."),
     ("Outras despesas sem nota/sem crédito (aluguel PF, tarifas, etc.)",4000,BRL,""),
     ("Regime tributário atual","Lucro Presumido",None,"Lista: MEI, Simples Nacional, Lucro Presumido, Lucro Real"),
     ("Impostos sobre a venda HOJE (soma, % do faturamento)",0.1165,PCT,"Presumido: PIS/Cofins 3,65% + ICMS/ISS efetivo. Simples: alíquota efetiva do DAS. Pergunte ao contador."),
     ("Desses, quanto é PIS/Cofins (% do faturamento)",0.0365,PCT,"Presumido 3,65%; Real até 9,25%; Simples: parte de PIS/Cofins do DAS."),
     ("Crédito de PIS/Cofins que você aproveita hoje (% das compras)",0,PCT,"Presumido: 0. Real: em geral 9,25%."),
     ("Crédito de ICMS que você aproveita hoje (% das compras)",0,PCT,"Comércio/indústria no regime normal de ICMS. Serviços e Simples: 0."),
     ("Redução de alíquota do seu item principal",0,PCT1,"0% (cheia), 30% (profissões regulamentadas), 60% (saúde, educação, alimentos listados) ou 100% (alíquota zero)."),
     ("% das vendas para outras empresas",0.3,PCT1,"Usado na aba 4.")]
    hdr(ws,4,["Item","Valor","Ajuda"])
    for i,(a,v,f,h) in enumerate(rows):
        r=5+i;lab(ws.cell(r,1,a));inp(ws.cell(r,2,v),f);ws.cell(r,3,h).font=Font(size=9,color="5B6576");ws.cell(r,3).alignment=Alignment(wrap_text=True);ws.row_dimensions[r].height=32
    dv=DataValidation(type="list",formula1='"MEI,Simples Nacional,Lucro Presumido,Lucro Real"',allow_blank=False);ws.add_data_validation(dv);dv.add("B10")
    dv2=DataValidation(type="list",formula1='"0,0.3,0.6,1"');ws.add_data_validation(dv2);dv2.add("B15")
    widths(ws,{"A":58,"B":20,"C":70})
    # 2 aliquotas
    wa=wb.create_sheet("2. Alíquotas");title(wa,"2. Alíquotas (EDITÁVEIS)","Estimativas não oficiais em 28/10/2026 · troque quando sair a alíquota oficial",6)
    lab(wa.cell(4,1,"CBS (alíquota de referência)"),True);inp(wa.cell(4,2,0.0921),PCT);wa.cell(4,3,"Estimativa não oficial do Comitê Gestor. Oficial: Senado fixa até 15/12/2026.").font=Font(size=9,color="5B6576")
    lab(wa.cell(5,1,"IBS no modelo final (estados + municípios)"),True);inp(wa.cell(5,2,0.187),PCT);wa.cell(5,3,"Estimativa não oficial. Total estimado 27,91%; teto de referência da lei: 26,5% em 2033, com ajuste.").font=Font(size=9,color="5B6576")
    lab(wa.cell(6,1,"Alíquota total no modelo final"),True);res(wa.cell(6,2,"=B4+B5"),PCT)
    hdr(wa,8,["Ano","CBS","IBS","PIS/Cofins que ainda existe","ICMS/ISS que ainda existe","Observação"])
    sched=[(2026,"0","0",1,1,"Ano de teste: CBS 0,9% e IBS 0,1% só destacados (compensáveis). Mantido 0 aqui."),
           (2027,"=B4","0.001",0,1,"CBS substitui PIS/Cofins. IBS 0,1%. IPI zera (exceto ZFM)."),
           (2028,"=B4","0.001",0,1,"Igual a 2027."),
           (2029,"=B4","=B5*0.1",0,0.9,"ICMS/ISS a 90%."),(2030,"=B4","=B5*0.2",0,0.8,"ICMS/ISS a 80%."),
           (2031,"=B4","=B5*0.3",0,0.7,"ICMS/ISS a 70%."),(2032,"=B4","=B5*0.4",0,0.6,"ICMS/ISS a 60%."),(2033,"=B4","=B5",0,0,"Modelo completo.")]
    for i,(y,c,ib,p,ic,o) in enumerate(sched):
        r=9+i;lab(wa.cell(r,1,y),True)
        for j,v in enumerate([c,ib,p,ic]):
            cell=wa.cell(r,2+j,float(v) if isinstance(v,str) and not v.startswith("=") else v);inp(cell,PCT)
        wa.cell(r,6,o).font=Font(size=9,color="5B6576")
    wa.cell(18,1,"Proporção de IBS de 2029 a 2032: aproximação didática (10%, 20%, 30%, 40% do IBS final). A proporção exata depende da regulamentação. Edite se o seu contador indicar outro valor.").font=Font(italic=True,size=9)
    widths(wa,{"A":42,"B":14,"C":14,"D":22,"E":22,"F":70})
    # 3 resultado
    wr=wb.create_sheet("3. Resultado");title(wr,"3. Resultado ano a ano","Regime regular de IBS/CBS (Presumido/Real ou Simples híbrido). Simples 'por dentro': veja a aba 4.",10)
    D="'1. Seus dados'!";A="'2. Alíquotas'!"
    hdr(wr,4,["Ano","Alíquota IBS+CBS efetiva","Débito IBS/CBS (por fora)","PIS/Cofins + ICMS/ISS","(−) Créditos","Imposto líquido","Carga sobre o faturamento","Variação x hoje","Lucro do mês (mesmo preço)","Reajuste de preço p/ manter o lucro"])
    wr.row_dimensions[4].height=45
    # hoje
    r=5;lab(wr.cell(r,1,"Hoje"),True)
    res(wr.cell(r,2,0),PCT);res(wr.cell(r,3,0),BRL)
    res(wr.cell(r,4,f"={D}B5*{D}B11"),BRL)
    res(wr.cell(r,5,f"={D}B6*({D}B13+{D}B14)*{D}B7"),BRL)
    res(wr.cell(r,6,"=C5+D5-E5"),BRL);res(wr.cell(r,7,f"=F5/{D}B5"),PCT);res(wr.cell(r,8,"=0"),PCT)
    res(wr.cell(r,9,f"={D}B5-F5-{D}B6-{D}B8-{D}B9"),BRL);res(wr.cell(r,10,0),PCT)
    for i in range(8):
        r=6+i;ar=9+i
        lab(wr.cell(r,1,f"={A}A{ar}"),True)
        res(wr.cell(r,2,f"=({A}B{ar}+{A}C{ar})*(1-{D}B15)"),PCT)
        res(wr.cell(r,3,f"={D}B5/(1+B{r})*B{r}"),BRL)
        res(wr.cell(r,4,f"={D}B5*({D}B12*{A}D{ar}+({D}B11-{D}B12)*{A}E{ar})"),BRL)
        res(wr.cell(r,5,f"={D}B6*{D}B7*(({A}B{ar}+{A}C{ar})/(1+{A}B{ar}+{A}C{ar})+{D}B13*{A}D{ar}+{D}B14*{A}E{ar})"),BRL)
        res(wr.cell(r,6,f"=MAX(0,C{r}+D{r}-E{r})"),BRL)
        res(wr.cell(r,7,f"=F{r}/{D}B5"),PCT)
        res(wr.cell(r,8,f"=G{r}-$G$5"),PCT)
        res(wr.cell(r,9,f"={D}B5-F{r}-{D}B6-{D}B8-{D}B9"),BRL)
        res(wr.cell(r,10,f"=IFERROR(($I$5-I{r})/(1/(1+B{r})-({D}B11-{D}B12)*{A}E{ar})/{D}B5,0)"),PCT)
    wr.conditional_formatting.add("H6:H13",CellIsRule(operator="greaterThan",formula=["0.005"],fill=RED))
    wr.conditional_formatting.add("H6:H13",CellIsRule(operator="lessThan",formula=["-0.005"],fill=GRN))
    wr.conditional_formatting.add("J6:J13",CellIsRule(operator="greaterThan",formula=["0.005"],fill=RED))
    wr.cell(15,1,"Como ler: 'Débito por fora' considera que o seu faturamento atual já inclui o IBS/CBS (mesmo preço). 'Reajuste' é quanto o preço precisa subir (ou pode cair, se negativo) para o lucro do mês ficar igual ao de hoje.").alignment=Alignment(wrap_text=True)
    wr.merge_cells("A15:J15");wr.row_dimensions[15].height=36
    wr.cell(16,1,"Estimativa educativa. Confirme com o seu contador antes de mudar preços.").font=Font(italic=True,color="B91C1C")
    widths(wr,{"A":10,"B":14,"C":17,"D":17,"E":15,"F":16,"G":14,"H":13,"I":18,"J":18})
    ch=BarChart();ch.type="col";ch.title="Carga de impostos sobre o faturamento";ch.y_axis.numFmt="0%";ch.height=7.5;ch.width=18
    ch.add_data(Reference(wr,min_col=7,min_row=4,max_row=13),titles_from_data=True);ch.set_categories(Reference(wr,min_col=1,min_row=5,max_row=13));ch.legend=None
    wr.add_chart(ch,"B18")
    # 4 simples
    wsx=wb.create_sheet("4. Simples x Híbrido");title(wsx,"4. Simples: por dentro x híbrido","Para empresas do Simples Nacional · referência: modelo final (2033)",4)
    items=[("Faturamento do mês",f"={D}B5",BRL,False),("Alíquota efetiva do DAS hoje",0.08,PCT,True),("Parte de IBS/CBS dentro do DAS (% da venda) — peça ao contador",0.03,PCT,True),
     ("% das vendas para empresas",f"={D}B16",PCT1,False),("Compras com nota de fornecedores do regime regular",f"={D}B6*{D}B7",BRL,False),
     ("Alíquota IBS+CBS (modelo final)",f"={A}B6*(1-{D}B15)",PCT,False),("Custo extra estimado de contabilidade/sistema no híbrido (mês)",300,BRL,True)]
    hdr(wsx,4,["Dado","Valor"])
    for i,(a,v,f,isin) in enumerate(items):
        r=5+i;lab(wsx.cell(r,1,a));(inp if isin else res)(wsx.cell(r,2,v),f)
    hdr(wsx,13,["Resultado","Por dentro","Híbrido","Diferença"])
    calc=[("Imposto do mês (DAS + IBS/CBS − créditos)","=B5*B6","=B5*(B6-B7)+B5/(1+B10)*B10-B9*B10/(1+B10)+B11"),
          ("Crédito que o cliente empresa aproveita a cada R$ 1.000","=1000*B7","=1000/(1+B10)*B10"),
          ("Crédito total gerado aos clientes empresa no mês","=B5*B8*B7","=B5*B8/(1+B10)*B10")]
    for i,(a,f1,f2) in enumerate(calc):
        r=14+i;lab(wsx.cell(r,1,a));res(wsx.cell(r,2,f1),BRL);res(wsx.cell(r,3,f2),BRL);res(wsx.cell(r,4,f"=C{r}-B{r}"),BRL)
    wsx.cell(18,1,"Leitura:").font=Font(bold=True)
    wsx.cell(19,1,'=IF(B8<0.2,"Você vende pouco para empresas: ficar POR DENTRO tende a ser mais simples e barato. Confirme com o contador.",IF(D14<=0,"O híbrido sai mais barato no imposto e ainda gera mais crédito ao cliente: vale avaliar com o contador.","O híbrido custa mais imposto por mês (célula D14), mas gera mais crédito aos seus clientes empresa (célula D16). Vale se os clientes PJ pagariam mais ou deixariam de comprar de você por falta de crédito. Decida com o contador."))')
    wsx.merge_cells("A19:D19");wsx.cell(19,1).alignment=Alignment(wrap_text=True);wsx.row_dimensions[19].height=48;wsx.cell(19,1).fill=GREY
    wsx.cell(21,1,"Prazos: a 1ª janela (efeito em jan/2027) encerrou em 30/10/2026. Próximas: março/2027 (efeito em 01/07/2027) e depois março e setembro de cada ano.").font=Font(italic=True,size=9)
    widths(wsx,{"A":62,"B":18,"C":18,"D":18})
    p=OUT+"Calculadora-Impacto-Reforma.xlsx";wb.save(p);return p

# ---------------- FLUXO 13 SEMANAS + DRE ----------------
def fluxo_sheet(wb,name="Fluxo 13 semanas",cfg="'Configurar'!"):
    ws=wb.create_sheet(name);title(ws,"Fluxo de caixa · 13 semanas","Preencha entradas e saídas previstas. Toda segunda: troque previsto por realizado.",16)
    ws.cell(4,1,"Semana").font=Font(bold=True)
    hdr(ws,4,["Semana"]+[f"S{i}" for i in range(1,14)]+["Total"])
    lab(ws.cell(5,1,"Começa em"),True)
    for i in range(13):
        c=ws.cell(5,2+i,f"={cfg}B5+{7*i}");c.number_format="dd/mm";res(c,None,False)
    ent=["Vendas à vista / Pix","Cartão (o que cai na conta)","Boletos / duplicatas a receber","Outras entradas"]
    sai=["Fornecedores / mercadoria","Folha + encargos","Pró-labore","Aluguel e condomínio","Água, luz, internet, sistemas","Impostos (DAS, guias)","Parcelas de empréstimos","Marketing","Outras saídas"]
    r=7;ws.cell(r-1,1,"ENTRADAS").font=Font(bold=True,color="15803D")
    for e in ent:
        lab(ws.cell(r,1,e))
        for i in range(13): inp(ws.cell(r,2+i,0),BRL0)
        res(ws.cell(r,15,f"=SUM(B{r}:N{r})"),BRL0);r+=1
    te=r;lab(ws.cell(r,1,"Total de entradas"),True)
    for i in range(14):
        col=ws.cell(r,2+i).column_letter;res(ws.cell(r,2+i,f"=SUM({col}7:{col}{r-1})"),BRL0)
    r+=2;ws.cell(r-1,1,"SAÍDAS").font=Font(bold=True,color="B91C1C");s0=r
    for e in sai:
        lab(ws.cell(r,1,e))
        for i in range(13): inp(ws.cell(r,2+i,0),BRL0)
        res(ws.cell(r,15,f"=SUM(B{r}:N{r})"),BRL0);r+=1
    ts=r;lab(ws.cell(r,1,"Total de saídas"),True)
    for i in range(14):
        col=ws.cell(r,2+i).column_letter;res(ws.cell(r,2+i,f"=SUM({col}{s0}:{col}{r-1})"),BRL0)
    r+=2
    lab(ws.cell(r,1,"Saldo inicial da semana"),True);si=r
    for i in range(13):
        col=ws.cell(r,2+i).column_letter;prev=ws.cell(r+2,1+i).column_letter
        res(ws.cell(r,2+i,f"={cfg}B6" if i==0 else f"={prev}{r+2}"),BRL0)
    lab(ws.cell(r+1,1,"Resultado da semana"),True)
    for i in range(13):
        col=ws.cell(r,2+i).column_letter;res(ws.cell(r+1,2+i,f"={col}{te}-{col}{ts}"),BRL0)
    lab(ws.cell(r+2,1,"SALDO FINAL"),True)
    for i in range(13):
        col=ws.cell(r,2+i).column_letter;c=ws.cell(r+2,2+i,f"={col}{r}+{col}{r+1}");res(c,BRL0)
    sf=r+2
    lab(ws.cell(r+3,1,"Acima da reserva mínima?"),True)
    for i in range(13):
        col=ws.cell(r,2+i).column_letter;res(ws.cell(r+3,2+i,f'=IF({col}{sf}<0,"NEGATIVO",IF({col}{sf}<{cfg}B7,"Abaixo","OK"))'),None,False)
    ws.conditional_formatting.add(f"B{sf}:N{sf}",CellIsRule(operator="lessThan",formula=["0"],fill=RED,font=Font(bold=True,color="B91C1C")))
    ws.conditional_formatting.add(f"B{sf+1}:N{sf+1}",CellIsRule(operator="equal",formula=['"NEGATIVO"'],fill=RED))
    ws.conditional_formatting.add(f"B{sf+1}:N{sf+1}",CellIsRule(operator="equal",formula=['"Abaixo"'],fill=PatternFill("solid",fgColor="FEF3C7")))
    ws.conditional_formatting.add(f"B{sf+1}:N{sf+1}",CellIsRule(operator="equal",formula=['"OK"'],fill=GRN))
    lab(ws.cell(sf+3,1,"Menor saldo previsto"),True);res(ws.cell(sf+3,2,f"=MIN(B{sf}:N{sf})"),BRL0)
    lab(ws.cell(sf+4,1,"Primeira semana negativa"),True);res(ws.cell(sf+4,2,f'=IFERROR("S"&MATCH(TRUE,INDEX(B{sf}:N{sf}<0,0),0),"Nenhuma")'),None)
    ws.freeze_panes="B6";widths(ws,{"A":34,**{ws.cell(1,2+i).column_letter:12 for i in range(14)}})
    ch=BarChart();ch.type="col";ch.title="Saldo final por semana";ch.height=7;ch.width=24;ch.legend=None
    ch.add_data(Reference(ws,min_col=2,max_col=14,min_row=sf),from_rows=True);ch.set_categories(Reference(ws,min_col=2,max_col=14,min_row=4))
    ws.add_chart(ch,f"A{sf+7}")
    return ws,sf

def caixa():
    wb=Workbook();wb.remove(wb.active)
    leia(wb,"Fluxo de Caixa 13 Semanas + DRE",["Para que serve:","• Fluxo 13 semanas: enxergar o caixa 3 meses à frente e descobrir com antecedência a semana em que o dinheiro ficaria negativo.","• DRE mensal: saber se a empresa dá lucro de verdade, qual a margem de contribuição e o ponto de equilíbrio.",
      "Como usar:","• Aba Configurar: data de início (uma segunda-feira), saldo inicial e reserva mínima.","• Aba Fluxo 13 semanas: preencha as células amarelas. Toda segunda, troque o previsto pelo realizado.","• Aba DRE mensal: uma coluna por mês.","Funciona no Excel e no Google Planilhas."])
    wc=wb.create_sheet("Configurar");title(wc,"Configurar","",3)
    for i,(a,v,f) in enumerate([("Nome da empresa","Minha Empresa",None),("Segunda-feira de início da semana 1",date(2026,11,2),"dd/mm/yyyy"),("Saldo de caixa inicial (todas as contas da empresa)",15000,BRL),("Reserva mínima desejada (1 mês de custo fixo)",18000,BRL)]):
        lab(wc.cell(4+i,1,a),True);inp(wc.cell(4+i,2,v),f)
    widths(wc,{"A":52,"B":22})
    fluxo_sheet(wb)
    wd=wb.create_sheet("DRE mensal");title(wd,"DRE simples · mês a mês","Regime de competência: o que foi vendido e gasto no mês",15)
    meses=["Jan","Fev","Mar","Abr","Mai","Jun","Jul","Ago","Set","Out","Nov","Dez"]
    hdr(wd,4,["Linha"]+meses+["Ano"])
    L=[("(+) Faturamento bruto","in"),("(−) Impostos sobre a venda","in"),("(=) Receita líquida","=B5-B6"),("(−) Custo da mercadoria / insumos / serviço","in"),("(−) Taxas de cartão, comissões, fretes de venda","in"),
       ("(=) Margem de contribuição","=B7-B8-B9"),("Margem de contribuição %","=IFERROR(B10/B5,0)"),("(−) Folha + encargos","in"),("(−) Pró-labore","in"),("(−) Aluguel, contas, sistemas","in"),("(−) Outras despesas fixas","in"),
       ("(=) Total de despesas fixas","=SUM(B12:B15)"),("(=) Resultado operacional","=B10-B16"),("(−) Juros e parcelas de empréstimos (juros)","in"),("(=) LUCRO LÍQUIDO","=B17-B18"),("Lucro líquido %","=IFERROR(B19/B5,0)"),("Ponto de equilíbrio (faturamento)","=IFERROR((B16+B18)/B11,0)")]
    for i,(a,f) in enumerate(L):
        r=5+i;lab(wd.cell(r,1,a),a.startswith("(=)") or a.startswith("Ponto"))
        for m in range(13):
            col=wd.cell(r,2+m).column_letter
            if m==12:
                if "%" in a: res(wd.cell(r,14,f"=IFERROR({ 'N10' if i==6 else 'N19'}/N5,0)"),PCT1)
                elif a.startswith("Ponto"): res(wd.cell(r,14,"=IFERROR((N16+N18)/N11,0)"),BRL0)
                else: res(wd.cell(r,14,f"=SUM(B{r}:M{r})"),BRL0)
                continue
            if f=="in": inp(wd.cell(r,2+m,0),BRL0)
            else:
                ff=f.replace("B",col) if col!="B" else f
                import re
                ff=re.sub(r"\bB(\d+)",lambda mm:col+mm.group(1),f)
                res(wd.cell(r,2+m,ff),PCT1 if "%" in a else BRL0)
    wd.conditional_formatting.add("B19:N19",CellIsRule(operator="lessThan",formula=["0"],fill=RED))
    wd.freeze_panes="B5";widths(wd,{"A":46,**{wd.cell(1,2+i).column_letter:12 for i in range(13)}})
    p=OUT+"Fluxo-de-Caixa-13-Semanas-e-DRE.xlsx";wb.save(p);return p

# ---------------- PAINEL DE IMPLEMENTAÇÃO ----------------
def painel():
    wb=Workbook();wb.remove(wb.active)
    wi=wb.create_sheet("Início");title(wi,"Painel de Implementação 2027","Kit Implementação 2027 · Kit Não Vire Estatística",5)
    for i,(a,v,f) in enumerate([("Empresa","Minha Empresa",None),("Segmento","Comércio e Alimentação",None),("Data de início do plano",date(2026,11,2),"dd/mm/yyyy"),("Saldo de caixa inicial",15000,BRL),("Reserva mínima (1 mês de custo fixo)",18000,BRL)]):
        lab(wi.cell(4+i,1,a),True);inp(wi.cell(4+i,2,v),f)
    dv=DataValidation(type="list",formula1='"Comércio e Alimentação,Serviços,Indústria / Agro / B2B"');wi.add_data_validation(dv);dv.add("B5")
    hdr(wi,10,["Resumo","Valor"])
    for i,(a,f,fm) in enumerate([("Tarefas do plano concluídas","=COUNTIF('Plano 30 dias'!F5:F34,\"OK\")&\" de 30\"",None),("Progresso do plano","=COUNTIF('Plano 30 dias'!F5:F34,\"OK\")/30",PCT1),
        ("Última nota do diagnóstico","=IFERROR(LOOKUP(2,1/('Nota mês a mês'!B5:B28<>\"\"),'Nota mês a mês'!B5:B28),\"—\")",None),
        ("Menor saldo previsto em 13 semanas","='Caixa semanal'!B"+"{SF3}",BRL0),("Primeira semana negativa","='Caixa semanal'!B"+"{SF4}",None),
        ("Próxima data da reforma","=IFERROR(INDEX(Calendário!B5:B16,MATCH(TRUE,INDEX(Calendário!A5:A16>=TODAY(),0),0)),\"—\")",None),
        ("Faltam (dias)","=IFERROR(INDEX(Calendário!A5:A16,MATCH(TRUE,INDEX(Calendário!A5:A16>=TODAY(),0),0))-TODAY(),\"—\")",None)]):
        lab(wi.cell(11+i,1,a),True);res(wi.cell(11+i,2,f),fm)
    wi.conditional_formatting.add("B12",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color="22C55E"))
    widths(wi,{"A":42,"B":34,"C":12,"D":12,"E":12})
    # plano
    wp=wb.create_sheet("Plano 30 dias");title(wp,"Plano 30 dias","Marque OK na coluna Status. As datas saem da aba Início.",6)
    hdr(wp,4,["Dia","Data","Escudo","Tarefa","Como fazer","Status"])
    for i,t in enumerate(NVE["plano30"]):
        r=5+i;wp.cell(r,1,t["d"]);c=wp.cell(r,2,f"=Início!$B$6+{t['d']-1}");c.number_format="dd/mm/yyyy"
        wp.cell(r,3,{"caixa":"Caixa","margem":"Margem","fiscal":"Fiscal"}[t["e"]])
        tt=t.get("t") or t["tN"]["A"]; cc=t.get("c") or t["cN"]["A"]
        if "tN" in t or "cN" in t:
            tt=t.get("t") or " | ".join(f"{k}: {v}" for k,v in t["tN"].items()); cc=t.get("c") or " | ".join(f"{ {'A':'Comércio','B':'Serviços','C':'Indústria/B2B'}[k]}: {v}" for k,v in t["cN"].items())
        wp.cell(r,4,tt);wp.cell(r,5,cc);inp(wp.cell(r,6,""))
        for col in range(1,6): wp.cell(r,col).alignment=Alignment(wrap_text=True,vertical="top");wp.cell(r,col).border=BOX
        wp.row_dimensions[r].height=max(30,14*(len(cc)//70+1))
    dv=DataValidation(type="list",formula1='"OK,Em andamento,Pendente"');wp.add_data_validation(dv);dv.add("F5:F34")
    wp.conditional_formatting.add("A5:F34",FormulaRule(formula=['$F5="OK"'],fill=GRN))
    wp.conditional_formatting.add("A5:F34",FormulaRule(formula=['AND($B5=TODAY(),$F5<>"OK")'],fill=PatternFill("solid",fgColor="FEF3C7")))
    widths(wp,{"A":6,"B":12,"C":10,"D":40,"E":70,"F":15});wp.freeze_panes="A5"
    # caixa semanal
    ws,sf=fluxo_sheet(wb,"Caixa semanal","Início!")
    # fix config refs: Início B6=data, B7=saldo, B8=reserva
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and "Início!" in c.value:
                c.value=c.value.replace("Início!B7","Início!__R").replace("Início!B6","Início!B7").replace("Início!B5","Início!B6").replace("Início!__R","Início!B8")
    for r_ in (14,15):
        cell=wi.cell(r_,2);cell.value=cell.value.replace("{SF3}",str(sf+3)).replace("{SF4}",str(sf+4))
    # nota
    wn=wb.create_sheet("Nota mês a mês");title(wn,"Nota mês a mês","Refaça o diagnóstico no app 1 vez por mês e copie as notas",6)
    hdr(wn,4,["Mês","Nota (0–100)","Caixa %","Margem %","Fiscal %","Faixa"])
    for i in range(24):
        r=5+i;c=wn.cell(r,1,f"=DATE(YEAR(Início!$B$6),MONTH(Início!$B$6)+{i},1)");c.number_format="mmm/yy"
        for col in range(2,6): inp(wn.cell(r,col,None))
        res(wn.cell(r,6,f'=IF(B{r}="","",IF(B{r}<40,"Zona de risco",IF(B{r}<70,"Atenção","Blindado em construção")))'),None,False)
    wn.conditional_formatting.add("B5:B28",CellIsRule(operator="between",formula=["0","39"],fill=RED))
    wn.conditional_formatting.add("B5:B28",CellIsRule(operator="between",formula=["70","100"],fill=GRN))
    lc=LineChart();lc.title="Evolução da nota";lc.height=8;lc.width=20;lc.y_axis.scaling.min=0;lc.y_axis.scaling.max=100
    lc.add_data(Reference(wn,min_col=2,max_col=5,min_row=4,max_row=16),titles_from_data=True);lc.set_categories(Reference(wn,min_col=1,min_row=5,max_row=16))
    wn.add_chart(lc,"H4");widths(wn,{"A":10,"B":14,"C":11,"D":11,"E":11,"F":22})
    # calendario
    wc=wb.create_sheet("Calendário");title(wc,"Calendário da reforma","Datas oficiais e o que fazer · coluna 'Faltam' conta os dias",6)
    hdr(wc,4,["Data","Evento","O que acontece","O que fazer (Comércio)","O que fazer (Serviços)","O que fazer (Indústria/Agro/B2B)","Faltam (dias)"])
    from datetime import datetime
    def dt(s):
        y,m,d=map(int,s.split("-"));return date(y,m,d)
    for i,x in enumerate(NVE["linha"]):
        r=5+i;c=wc.cell(r,1,dt(x["d"]));c.number_format="dd/mm/yyyy"
        wc.cell(r,2,x["t"]);wc.cell(r,3,x["txt"])
        for j,n in enumerate("ABC"): wc.cell(r,4+j,pick(x["acao"],n))
        res(wc.cell(r,7,f'=IF(A{r}<TODAY(),"já passou",A{r}-TODAY())'),None,False)
        for col in range(1,8): wc.cell(r,col).alignment=Alignment(wrap_text=True,vertical="top");wc.cell(r,col).border=BOX
        wc.row_dimensions[r].height=62
    wc.conditional_formatting.add("A5:G16",FormulaRule(formula=['AND($A5>=TODAY(),$A5-TODAY()<=45)'],fill=PatternFill("solid",fgColor="FEF3C7")))
    widths(wc,{"A":12,"B":30,"C":48,"D":34,"E":34,"F":34,"G":12});wc.freeze_panes="A5"
    wc.cell(18,1,"Fontes: EC 132/2023; LC 214/2025; Res. CGSN 194/2026; Receita Federal; Comsefaz; Senado. Dados de 28/10/2026.").font=Font(italic=True,size=9)
    # simples x hibrido
    wsx=wb.create_sheet("Simples x Híbrido");title(wsx,"Simples: por dentro x híbrido","Faça antes das janelas de março e setembro · valores do mês",4)
    items=[("Faturamento do mês",40000,BRL),("Alíquota efetiva do DAS hoje",0.08,PCT),("Parte de IBS/CBS dentro do DAS (% da venda) — peça ao contador",0.03,PCT),("% das vendas para empresas",0.3,PCT1),
           ("Compras com nota de fornecedores do regime regular",12000,BRL),("Alíquota IBS+CBS (modelo final, estimativa)",0.2791,PCT),("Redução de alíquota do seu item (0, 30%, 60%)",0,PCT1),("Custo extra de contabilidade/sistema no híbrido",300,BRL)]
    hdr(wsx,4,["Dado","Valor"])
    for i,(a,v,f) in enumerate(items): lab(wsx.cell(5+i,1,a));inp(wsx.cell(5+i,2,v),f)
    lab(wsx.cell(13,1,"Alíquota efetiva no seu item"),True);res(wsx.cell(13,2,"=B10*(1-B11)"),PCT)
    hdr(wsx,15,["Resultado","Por dentro","Híbrido","Diferença"])
    calc=[("Imposto do mês","=B5*B6","=B5*(B6-B7)+B5/(1+B13)*B13-B9*B10/(1+B10)+B12"),("Crédito do cliente PJ a cada R$ 1.000","=1000*B7","=1000/(1+B13)*B13"),("Crédito gerado aos clientes PJ no mês","=B5*B8*B7","=B5*B8/(1+B13)*B13")]
    for i,(a,f1,f2) in enumerate(calc):
        r=16+i;lab(wsx.cell(r,1,a));res(wsx.cell(r,2,f1),BRL);res(wsx.cell(r,3,f2),BRL);res(wsx.cell(r,4,f"=C{r}-B{r}"),BRL)
    wsx.cell(20,1,'=IF(B8<0.2,"Você vende pouco para empresas: ficar POR DENTRO tende a ser o caminho. Confirme com o contador.",IF(D16<=0,"O híbrido sai mais barato e gera mais crédito ao cliente: avalie com o contador.","O híbrido custa mais imposto por mês (célula D16), mas gera mais crédito aos clientes PJ (célula D18). Vale se os clientes PJ pagariam mais ou deixariam de comprar de você por falta de crédito. Decida com o contador."))')
    wsx.merge_cells("A20:D20");wsx.cell(20,1).alignment=Alignment(wrap_text=True);wsx.row_dimensions[20].height=44;wsx.cell(20,1).fill=GREY
    widths(wsx,{"A":62,"B":18,"C":18,"D":18})
    # preco
    wpr=wb.create_sheet("Preço 2027");title(wpr,"Preço 2027 · item a item","Regime regular (Presumido/Real/híbrido). Alíquotas editáveis nas células amarelas do topo.",13)
    lab(wpr.cell(4,1,"CBS (estimativa; oficial até 15/12/2026)"),True);inp(wpr.cell(4,2,0.0921),PCT)
    lab(wpr.cell(5,1,"IBS final (estimativa)"),True);inp(wpr.cell(5,2,0.187),PCT)
    lab(wpr.cell(6,1,"PIS/Cofins hoje (% da venda)"),True);inp(wpr.cell(6,2,0.0365),PCT)
    lab(wpr.cell(7,1,"% das compras com crédito cheio"),True);inp(wpr.cell(7,2,0.7),PCT1)
    hdr(wpr,9,["Item","Preço hoje","Custo direto","Impostos hoje %","Taxas/comissão %","Crédito hoje % (PIS+ICMS)","Crédito ICMS hoje %","Redução (0/0,3/0,6/1)","Lucro hoje","Preço 2027 p/ manter lucro","Variação 2027","Preço 2033 p/ manter lucro","Variação 2033"])
    wpr.row_dimensions[9].height=45
    for i in range(15):
        r=10+i
        if i==0: vals=["Produto/serviço exemplo",100,50,0.12,0.04,0,0,0]
        else: vals=["",None,None,None,None,None,None,None]
        wpr.cell(r,1,vals[0]);inp(wpr.cell(r,1))
        for j,fm in enumerate([BRL,BRL,PCT,PCT,PCT,PCT,PCT1]): inp(wpr.cell(r,2+j,vals[1+j]),fm)
        g=f"IF(B{r}=\"\",\"\","
        lucro=f"B{r}*(1-D{r}-E{r})-C{r}+C{r}*F{r}*$B$7"
        res(wpr.cell(r,9,f"={g}{lucro})"),BRL)
        c27=f"($B$4*(1-H{r})+IF(H{r}<1,0.001,0))";out=f"MAX(0,D{r}-$B$6)"
        cred27=f"(C{r}*$B$7*$B$4/(1+$B$4+{out})+C{r}*$B$7*G{r})"
        kA=f"((1+{c27})*(1-{out}-E{r})-{c27})"
        res(wpr.cell(r,10,f"={g}(I{r}+C{r}-{cred27})/{kA}*(1+{c27}))"),BRL)
        res(wpr.cell(r,11,f"={g}J{r}/B{r}-1)"),PCT1)
        t=f"(($B$4+$B$5)*(1-H{r}))";tf="($B$4+$B$5)";cred33=f"(C{r}*$B$7*{tf}/(1+{tf}))"
        res(wpr.cell(r,12,f"={g}(I{r}+C{r}-{cred33})/(1-E{r}*(1+{t}))*(1+{t}))"),BRL)
        res(wpr.cell(r,13,f"={g}L{r}/B{r}-1)"),PCT1)
    wpr.conditional_formatting.add("K10:K24",CellIsRule(operator="greaterThan",formula=["0.005"],fill=RED))
    wpr.conditional_formatting.add("M10:M24",CellIsRule(operator="greaterThan",formula=["0.005"],fill=RED))
    wpr.cell(26,1,"Preço 2027/2033 já inclui o IBS/CBS 'por fora'. Estimativa educativa: confirme com o contador. Mesma conta do simulador Preço 2027 do app.").font=Font(italic=True,size=9)
    widths(wpr,{"A":30,"B":12,"C":12,"D":12,"E":12,"F":14,"G":12,"H":12,"I":12,"J":15,"K":11,"L":15,"M":11});wpr.freeze_panes="B10"
    p=OUT+"Painel-de-Implementacao-2027.xlsx";wb.save(p);return p

if __name__=="__main__":
    print(calculadora());print(caixa());print(painel())
